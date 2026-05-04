#!/usr/bin/env python3
"""
Konversi file Excel Aktivitas Kinerja Harian ke format JSON
untuk digunakan oleh userscript atau skrip Selenium e-MASTER.

Output JSON dikelompokkan per "Kegiatan Tugas Jabatan" (breakdown),
sehingga bisa langsung digunakan per halaman realisasi di e-MASTER.

Penggunaan:
    python convert_excel.py <file_excel.xlsx> [--output output.json]

Contoh:
    python convert_excel.py "Aktifitas Kinerja Harian.xlsx"
    python convert_excel.py data.xlsx --output april2026.json
"""

import argparse
import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl belum terinstall. Jalankan: pip install openpyxl")
    sys.exit(1)


def parse_excel(file_path: str) -> dict:
    """Parse file Excel Aktivitas Kinerja Harian dan konversi ke JSON."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active  # Gunakan sheet pertama

    # Cari header info (Nama, NIP, Unit Kerja)
    info = {"nama": "", "nip": "", "unit_kerja": "", "file_source": str(file_path)}

    for row_idx in range(1, min(20, ws.max_row + 1)):
        cell_a = ws.cell(row=row_idx, column=1).value
        cell_b = ws.cell(row=row_idx, column=2).value
        if cell_a and isinstance(cell_a, str):
            lower_a = cell_a.strip().lower()
            if lower_a == "nama" and cell_b:
                info["nama"] = str(cell_b).strip().lstrip(": ")
            elif lower_a == "nip" and cell_b:
                info["nip"] = str(cell_b).strip().lstrip(": ")
            elif "unit kerja" in lower_a and cell_b:
                info["unit_kerja"] = str(cell_b).strip().lstrip(": ")

    # Cari baris header tabel (No, Hari, Tanggal, Kegiatan, ...)
    header_row = None
    for row_idx in range(1, min(30, ws.max_row + 1)):
        cell_a = ws.cell(row=row_idx, column=1).value
        if cell_a and str(cell_a).strip().lower() == "no":
            header_row = row_idx
            break

    if header_row is None:
        print("Error: Header tabel (No, Hari, Tanggal, ...) tidak ditemukan!")
        sys.exit(1)

    # Baca header kolom
    headers = []
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col_idx).value
        headers.append(str(val).strip() if val else f"col_{col_idx}")

    print(f"Header ditemukan di baris {header_row}: {headers}")

    # Mapping kolom berdasarkan header
    col_map = {}
    for i, h in enumerate(headers):
        h_lower = h.lower()
        if h_lower == "no":
            col_map["no"] = i
        elif h_lower in ("hari",):
            col_map["hari"] = i
        elif "tanggal" in h_lower or "tgl" in h_lower:
            col_map["tanggal"] = i
        elif "kegiatan" in h_lower or "detail" in h_lower or "aktivitas" in h_lower:
            col_map["kegiatan"] = i
        elif "obyek" in h_lower or "objek" in h_lower or "uraian" in h_lower:
            col_map["objek_kerja"] = i
        elif "volume" in h_lower or "vol" in h_lower:
            col_map["volume"] = i
        elif "beban" in h_lower:
            col_map["beban_kerja"] = i
        elif "satuan" in h_lower:
            col_map["satuan"] = i
        elif "durasi" in h_lower or "wpt" in h_lower or (
            "menit" in h_lower and "beban" not in h_lower
        ):
            col_map["durasi"] = i

    print(f"Kolom terdeteksi: {col_map}")

    # Baca data dan kelompokkan per kegiatan (breakdown)
    grouped: dict[str, list] = defaultdict(list)
    current_hari = ""
    current_tanggal = ""
    data_start = header_row + 1

    # Skip sub-header rows (rows after main header that are also headers)
    for skip_row in range(data_start, data_start + 3):
        cell = ws.cell(row=skip_row, column=1).value
        if cell is None or (isinstance(cell, str) and not cell.strip().isdigit()):
            data_start = skip_row + 1
        else:
            break

    for row_idx in range(data_start, ws.max_row + 1):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            row_data.append(ws.cell(row=row_idx, column=col_idx).value)

        # Skip empty rows dan baris total
        no_val = row_data[col_map.get("no", 0)] if "no" in col_map else None
        if no_val is None:
            continue
        if isinstance(no_val, str) and no_val.strip().lower() in ("total", "nb", ""):
            continue
        if isinstance(no_val, str) and not no_val.strip().isdigit():
            continue

        # Ambil hari & tanggal (hanya diisi di baris pertama tiap hari)
        hari_val = row_data[col_map["hari"]] if "hari" in col_map else None
        tgl_val = row_data[col_map["tanggal"]] if "tanggal" in col_map else None

        if hari_val:
            current_hari = str(hari_val).strip()
        if tgl_val:
            if isinstance(tgl_val, datetime):
                current_tanggal = tgl_val.strftime("%d-%m-%Y")
            else:
                current_tanggal = str(tgl_val).strip()

        # Ambil data kegiatan
        kegiatan = ""
        if "kegiatan" in col_map:
            val = row_data[col_map["kegiatan"]]
            kegiatan = str(val).strip() if val else ""

        objek_kerja = ""
        if "objek_kerja" in col_map:
            val = row_data[col_map["objek_kerja"]]
            objek_kerja = str(val).strip() if val else ""

        volume = ""
        if "volume" in col_map:
            val = row_data[col_map["volume"]]
            volume = str(int(val)) if val and isinstance(val, (int, float)) else str(val or "")

        durasi = ""
        if "durasi" in col_map:
            val = row_data[col_map["durasi"]]
            durasi = str(int(val)) if val and isinstance(val, (int, float)) else str(val or "")

        beban_kerja = ""
        if "beban_kerja" in col_map:
            val = row_data[col_map["beban_kerja"]]
            beban_kerja = str(int(val)) if val and isinstance(val, (int, float)) else str(val or "")

        satuan = "Pasien"
        if "satuan" in col_map:
            val = row_data[col_map["satuan"]]
            satuan = str(val).strip() if val else "Pasien"

        if not kegiatan:
            continue

        # Skip kegiatan non-breakdown (tidak ada di halaman realisasi e-MASTER)
        SKIP_KEGIATAN = {"briefing", "timbang terima"}
        if kegiatan.lower().strip() in SKIP_KEGIATAN:
            continue

        entry = {
            "hari": current_hari,
            "tanggal": current_tanggal,
            "detail_aktivitas": kegiatan,
            "objek_kerja": objek_kerja,
            "satuan": satuan,
            "wpt_menit": durasi,
            "volume": volume,
            "beban_kerja": beban_kerja,
        }

        # Kelompokkan berdasarkan nama kegiatan (breakdown)
        grouped[kegiatan].append(entry)

    wb.close()

    # Deteksi bulan dari tanggal pertama
    bulan = ""
    all_entries = []
    for entries in grouped.values():
        all_entries.extend(entries)
    if all_entries:
        first_date = all_entries[0].get("tanggal", "")
        parts = first_date.replace("/", "-").split("-")
        if len(parts) == 3:
            bulan = parts[1]  # DD-MM-YYYY -> MM

    # Kamus Aktifitas Harian - kata kunci pencarian per breakdown
    # Dipakai di popup "Detail Aktivitas" saat mengisi form Tambah Aktivitas
    KAMUS_KEYWORDS = {
        "Melaksanakan asuhan keperawatan sesuai SOP": {
            "keywords": ["Manajemen Asuhan Keperawatan"],
            "mode": "single",
        },
        "Menginput dokumentasi tindakan keperawatan": {
            "keywords": ["Memasukkan Hasil Pengkajian"],
            "mode": "single",
        },
        "Melaksanakan tindakan keperawatan tepat waktu": {
            "keywords": [
                "Sampling Darah Vena Instalasi",
                "Terapi Injeksi Parenteral",
                "Pasang Infus",
            ],
            "mode": "rotating",
        },
        "Melaksanakan prosedur keperawatan sesuai SOP": {
            # Urutan non-adjacent: 1,3,2,4 dari list asli
            "keywords": [
                "Terapi Injeksi Line",
                "Pasang Infus",
                "Mengukur Tanda Tanda",
                "Sampling Darah Vena Instalasi",
            ],
            "mode": "rotating",
        },
        "Menyiapkan alat medis pelayanan": {
            "keywords": ["Memeriksa Kelengkapan Alat"],
            "mode": "single",
        },
    }

    # Buat struktur breakdowns
    breakdowns = []
    total_expanded = 0
    for kegiatan_name, entries in grouped.items():
        kamus = KAMUS_KEYWORDS.get(kegiatan_name, {
            "keywords": [kegiatan_name],
            "mode": "single",
        })

        if kamus["mode"] == "rotating":
            # Expand: setiap entry (1 per tanggal) di-duplikasi jadi N entry,
            # satu untuk setiap kata kunci. Data tetap sama, hanya kamus_keyword beda.
            kw_list = kamus["keywords"]
            expanded_entries = []
            for entry in entries:
                for kw in kw_list:
                    new_entry = dict(entry)
                    new_entry["kamus_keyword"] = kw
                    expanded_entries.append(new_entry)
            entries = expanded_entries
        else:
            for entry in entries:
                entry["kamus_keyword"] = kamus["keywords"][0]

        total_expanded += len(entries)
        breakdowns.append({
            "kegiatan_tugas_jabatan": kegiatan_name,
            "jumlah_entries": len(entries),
            "kamus_config": kamus,
            "entries": entries,
        })

    result = {
        "info": info,
        "bulan": bulan,
        "generated_at": datetime.now().isoformat(),
        "total_entries": total_expanded,
        "total_breakdowns": len(breakdowns),
        "breakdowns": breakdowns,
    }

    return result


def main():
    parser = argparse.ArgumentParser(
        description="Konversi Excel Aktivitas Kinerja Harian ke JSON untuk e-MASTER"
    )
    parser.add_argument("excel_file", help="Path ke file Excel (.xlsx)")
    parser.add_argument(
        "--output", "-o",
        help="Path output file JSON (default: <nama_excel>.json)",
        default=None,
    )

    args = parser.parse_args()

    excel_path = Path(args.excel_file)
    if not excel_path.exists():
        print(f"Error: File '{excel_path}' tidak ditemukan!")
        sys.exit(1)

    output_path = args.output or str(excel_path.with_suffix(".json"))

    print(f"Membaca: {excel_path}")
    data = parse_excel(str(excel_path))

    print(f"\nHasil:")
    print(f"  Nama: {data['info']['nama']}")
    print(f"  NIP: {data['info']['nip']}")
    print(f"  Unit Kerja: {data['info']['unit_kerja']}")
    print(f"  Bulan: {data['bulan']}")
    print(f"  Total entry: {data['total_entries']}")
    print(f"  Jumlah breakdown: {data['total_breakdowns']}")
    print(f"\n  Breakdown per kegiatan:")
    for bd in data["breakdowns"]:
        print(f"    - {bd['kegiatan_tugas_jabatan']}: {bd['jumlah_entries']} entry")

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"\nJSON tersimpan di: {output_path}")


if __name__ == "__main__":
    main()
